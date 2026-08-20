"""
Report exporters — CSV and Excel (.xlsx).

Both exporters consume the same ``ReportData`` dict produced by
``analytics.report_data``. Every exported file carries a consistent professional
header block (system name, report title, generated date, applied filters and
summary statistics) followed by the detailed data table — so a CSV, an Excel
sheet and the on-screen/printed report all present identical information.

PDF export is handled by the print-friendly report page + the browser's native
"Save as PDF" (see the report templates' print CSS); no server-side PDF engine
is required.
"""

import csv

from django.http import HttpResponse
from django.utils.text import slugify


def _filename(data, ext):
    stamp = data["generated_at"].strftime("%Y%m%d-%H%M")
    return f"{slugify(data['title'])}-{stamp}.{ext}"


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def export_csv(data):
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{_filename(data, "csv")}"'
    writer = csv.writer(response)

    writer.writerow([data["system_name"]])
    writer.writerow([data["title"]])
    writer.writerow(["Generated", data["generated_at"].strftime("%Y-%m-%d %H:%M")])
    writer.writerow([])

    if data["filters"]:
        writer.writerow(["Applied Filters"])
        for label, value in data["filters"]:
            writer.writerow([label, value])
        writer.writerow([])

    if data["summary"]:
        writer.writerow(["Summary"])
        for card in data["summary"]:
            writer.writerow([card["label"], card["value"]])
        writer.writerow([])

    if data.get("note"):
        writer.writerow(["Note", data["note"]])
        writer.writerow([])

    writer.writerow(["Detailed Data"])
    writer.writerow(data["columns"])
    for row in data["rows"]:
        writer.writerow(row)

    return response


# ---------------------------------------------------------------------------
# Excel (.xlsx) via openpyxl
# ---------------------------------------------------------------------------
def export_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    brand = Font(bold=True, size=15, color="1F6F3D")
    subtitle = Font(bold=True, size=12, color="333333")
    label_font = Font(bold=True, color="555555")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6F3D")
    section_font = Font(bold=True, size=11, color="1F6F3D")
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ncols = max(len(data["columns"]), 2)

    def merge(row, text, font):
        ws.cell(row=row, column=1, value=text).font = font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)

    r = 1
    merge(r, data["system_name"], brand); r += 1
    merge(r, data["title"], subtitle); r += 1
    ws.cell(row=r, column=1, value="Generated").font = label_font
    ws.cell(row=r, column=2, value=data["generated_at"].strftime("%Y-%m-%d %H:%M"))
    r += 2

    if data["filters"]:
        merge(r, "Applied Filters", section_font); r += 1
        for lbl, val in data["filters"]:
            ws.cell(row=r, column=1, value=lbl).font = label_font
            ws.cell(row=r, column=2, value=str(val))
            r += 1
        r += 1

    if data["summary"]:
        merge(r, "Summary", section_font); r += 1
        for card in data["summary"]:
            ws.cell(row=r, column=1, value=card["label"]).font = label_font
            ws.cell(row=r, column=2, value=card["value"])
            r += 1
        r += 1

    if data.get("note"):
        ws.cell(row=r, column=1, value="Note").font = label_font
        note_cell = ws.cell(row=r, column=2, value=data["note"])
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncols)
        r += 2

    # Detail header
    merge(r, "Detailed Data", section_font); r += 1
    header_row = r
    for c, col in enumerate(data["columns"], start=1):
        cell = ws.cell(row=r, column=c, value=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    r += 1

    for row in data["rows"]:
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = border
        r += 1

    # Column widths from content
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        longest = 0
        for row_cells in ws.iter_rows(min_col=c, max_col=c):
            for cell in row_cells:
                if cell.value is not None:
                    longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 45)

    # Freeze detail header row
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(data, "xlsx")}"'
    wb.save(response)
    return response


def export_report(data, fmt):
    """Dispatch to the requested export format."""
    if fmt == "csv":
        return export_csv(data)
    if fmt in ("xlsx", "excel"):
        return export_xlsx(data)
    raise ValueError(f"Unsupported export format: {fmt}")
